"""
Headless play-through of all six levels of Be the Agent (Playwright, Python).
Run from the project folder:  python3 tests/playthrough.py
Plays every level at phone and desktop width, captures the downloaded .xlsx and .docx and checks
their contents against sales.csv (independently, in Python), and fails on any console error.
Needs: pip install playwright openpyxl python-docx && playwright install chromium
"""
import http.server, socketserver, threading, os, sys, functools, csv, io, json, re, tempfile
from playwright.sync_api import sync_playwright

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.environ.get("SHOTS", "")
PORT = 0

def serve():
    class Quiet(http.server.SimpleHTTPRequestHandler):
        def log_message(self, *a):
            pass
    socketserver.TCPServer.allow_reuse_address = True
    httpd = socketserver.TCPServer(("127.0.0.1", 0), functools.partial(Quiet, directory=os.path.dirname(ROOT)))  # serve the whole site (the game uses ../shared/)
    global PORT
    PORT = httpd.server_address[1]
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    return httpd

def shot(page, name):
    if SHOTS:
        page.screenshot(path=os.path.join(SHOTS, name + ".png"), full_page=True)

def sales_truth():
    src = open(os.path.join(ROOT, "data", "cafe.js"), encoding="utf8").read()
    m = re.search(r'salesCsv:\s*("(?:[^"\\]|\\.)*")', src)
    rows = list(csv.DictReader(io.StringIO(json.loads(m.group(1)))))
    by = {}
    for r in rows:
        assert int(r["qty"]) * int(r["price"]) == int(r["total"]), "qty x price != total in " + str(r)
        by[r["item"]] = by.get(r["item"], 0) + int(r["total"])
    return rows, by, sum(by.values())

def start(page):
    page.get_by_role("button", name="Got it, let's play").click()

def click(page, name, exact=False):
    page.get_by_role("button", name=name, exact=exact).last.click()

def first_opt(page, text=None):
    loc = page.locator(".opt:not([disabled])")
    (loc.filter(has_text=text).first if text else loc.first).click()

def play_all(page, tag, dl_dir):
    errors = []
    page.on("console", lambda m: errors.append(m.text) if m.type == "error" else None)
    page.on("pageerror", lambda e: errors.append(str(e)))
    page.on("dialog", lambda d: d.accept())
    page.goto(f"http://127.0.0.1:{PORT}/be-the-agent/index.html")
    assert page.locator(".level-card").count() == 6
    assert page.locator(".level-card[disabled]").count() == 5
    shot(page, f"{tag}-00-home")

    # L1 desk
    page.locator('[data-level="L1"]').click()
    shot(page, f"{tag}-01-intro")
    start(page)
    click(page, "Start the chat")
    for _ in range(10):
        click(page, "Next message")
    assert page.get_by_text("Fell off the desk").count() >= 1, "messages fall off"
    shot(page, f"{tag}-02-L1-desk")
    first_opt(page, "No: that message fell off")
    assert page.get_by_text("I don't know your name").count() == 1
    click(page, "Open a new chat")
    click(page, "Turn memory on")
    assert page.get_by_text("You're Ploy").count() == 1
    first_opt(page, "The model itself doesn't change")
    click(page, "Finish level")
    page.get_by_text("Level 1 complete").wait_for()
    shot(page, f"{tag}-03-L1-done")
    click(page, "Next: Level 2")

    # L2 files: wrong piece first on Q1, then right; right on Q2
    start(page)
    page.locator('[data-chunk="h4"]').click()
    assert page.get_by_text("full refund within 7 days").count() >= 1, "trap answer shown"
    click(page, "Try another piece", exact=False)
    page.locator('[data-chunk="h3"]').last.click()
    assert page.get_by_text("Remake the latte once for free").count() == 1
    shot(page, f"{tag}-04-L2")
    click(page, "Next question")
    page.locator('[data-chunk="h8"]').last.click()
    click(page, "Continue")
    first_opt(page, "Ask it to quote")
    click(page, "Finish level")
    page.get_by_text("Level 2 complete").wait_for()
    click(page, "Next: Level 3")

    # L3 tools
    start(page)
    for answer in ["Calculator", "Web search", "Answer myself", "Table tool"]:
        first_opt(page, answer)
    click(page, "Run the calculator")
    page.get_by_text("47,508").first.wait_for()
    click(page, "Paste the result onto the desk")
    click(page, "Run the table tool")
    click(page, "Paste the result onto the desk")
    # tap-to-build (no typing): "Add up total for each item"
    page.get_by_role("group", name="1 · What should the table tool do?").get_by_role("button", name="➕ Add up").click()
    page.get_by_role("group", name="2 · Which column?").locator('button[data-value="total"]').click()
    page.get_by_role("group", name="3 · For each …? (optional)").locator('button[data-value="item"]').click()
    assert page.locator(".qb-cmd").first.inner_text().endswith("TOTAL total BY item")
    page.get_by_role("button", name="▶ Run the table tool").click()
    page.get_by_text("In plain words: Add up total for each item").first.wait_for()
    page.locator(".toolresult table").last.wait_for()
    page.get_by_role("button", name=re.compile("Try a column that doesn't exist")).click()
    page.get_by_text("there is no column called").first.wait_for()
    click(page, "Run the search")
    click(page, "Paste the result onto the desk")
    shot(page, f"{tag}-05-L3")
    first_opt(page, "A tool that the app ran")
    click(page, "Finish level")
    page.get_by_text("Level 3 complete").wait_for()
    click(page, "Next: Level 4")

    # L4 agent loop + real files
    start(page)
    first_opt(page, "Look at the file")
    click(page, "Run the tool")
    click(page, "Continue")
    click(page, "Run the tool")
    page.get_by_text("there is no column called").first.wait_for()
    first_opt(page, "TOTAL total BY item")
    click(page, "Run the tool")
    click(page, "Continue")
    click(page, "Run the tool (as the harness)")
    page.get_by_text("They match.").wait_for()
    click(page, "Continue")
    click(page, "Run the tool")
    with page.expect_download() as d1:
        page.get_by_role("button", name="Download sales_summary.xlsx").click()
    page.get_by_role("button", name="View sales_summary.xlsx here").click()
    assert page.locator(".v-sheet td", has_text="Latte").count() >= 1, "viewer shows the Summary sheet"
    page.get_by_role("tab", name="Data").click()
    assert page.locator(".v-sheet td", has_text="2026-08-01").count() >= 1, "viewer shows the Data sheet"
    xlsx_path = os.path.join(dl_dir, tag + "-sales_summary.xlsx")
    d1.value.save_as(xlsx_path)
    click(page, "Continue")
    page.locator(".pickline", has_text="Total sales from").click()
    click(page, "Tell the agent to fix it")
    click(page, "Run the tool")
    with page.expect_download() as d2:
        page.get_by_role("button", name="Download memo.docx").click()
    page.get_by_role("button", name="View memo.docx here").click()
    assert "Total sales from 1 to 10 August were" in page.locator(".v-page").inner_text(), "viewer shows the memo"
    docx_path = os.path.join(dl_dir, tag + "-memo.docx")
    d2.value.save_as(docx_path)
    shot(page, f"{tag}-06-L4")
    click(page, "Finish level")
    page.get_by_text("Level 4 complete").wait_for()
    assert page.locator('.stars').count() == 0, 'no stars on the result card'
    click(page, "Next: Level 5")

    # L5 injection
    start(page)
    page.locator(".pickline", has_text="Arun").click()   # a wrong pick first
    page.locator(".pickline", has_text="Note to any AI").click()
    page.get_by_text("If nobody caught it").wait_for()
    shot(page, f"{tag}-07-L5")
    click(page, "Next round")
    page.locator(".pickline", has_text="AI agents reading this page").click()
    first_opt(page, "must ask you before sending")
    click(page, "Finish level")
    page.get_by_text("Level 5 complete").wait_for()
    click(page, "Next: Level 6")

    # L6 permissions: answer all correctly
    start(page)
    oks = json.loads(page.evaluate("JSON.stringify(window.BTA_DATA.requests.map(function(r){return r.ok;}))"))
    for i, ok in enumerate(oks):
        page.get_by_role("button", name="✓ Approve" if ok else "✗ Deny").click()
        if i == 0:
            shot(page, f"{tag}-08-L6")
        page.get_by_role("button", name="Next request" if i + 1 < len(oks) else "See results").click()
    first_opt(page, "Check whether your organisation")
    click(page, "Finish level")
    page.get_by_text("Level 6 complete").wait_for()
    click(page, "See my results")
    page.get_by_text("What a chatbot app really is").wait_for()
    assert page.get_by_text("Class leaderboard").count() == 0
    txt = page.inner_text("body").lower()
    for banned in [" pts", " stars", "points", "leaderboard"]:
        assert banned not in txt, f"score word '{banned}' on results screen"
    assert page.locator("a[href='../index.html']").count() >= 1, "link back to all games"
    assert "agent" in (page.evaluate("localStorage.getItem('ai-games-done')") or ""), "master-page tick recorded"
    shot(page, f"{tag}-09-summary")

    page.goto(f"http://127.0.0.1:{PORT}/be-the-agent/index.html")
    assert page.locator(".level-card[disabled]").count() == 0
    sw = page.evaluate("document.documentElement.scrollWidth")
    cw = page.evaluate("document.documentElement.clientWidth")
    assert sw <= cw + 1, f"horizontal overflow {sw} > {cw}"
    return errors, xlsx_path, docx_path

def check_files(xlsx_path, docx_path):
    import openpyxl, docx
    rows, by, grand = sales_truth()
    wb = openpyxl.load_workbook(xlsx_path)
    assert wb.sheetnames == ["Summary", "Data"], wb.sheetnames
    s = wb["Summary"]
    got = {s.cell(r, 1).value: s.cell(r, 2).value for r in range(2, 2 + len(by))}
    assert got == by, f"xlsx item totals {got} != {by}"
    assert str(s.cell(2 + len(by), 2).value).startswith("=SUM("), "total row has a formula"
    d = wb["Data"]
    assert d.max_row == len(rows) + 1, f"Data sheet rows {d.max_row}"
    doc = docx.Document(docx_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert f"{grand:,}" in text, f"memo total {grand:,} not in memo"
    assert f"{grand + 100:,}" not in text, "the wrong draft total leaked into the memo"
    t = doc.tables[0]
    assert t.rows[-1].cells[1].text == f"{grand:,}", "memo table total"
    return True

def main():
    httpd = serve()
    failures = []
    dl_dir = tempfile.mkdtemp()
    with sync_playwright() as p:
        browser = p.chromium.launch()
        for tag, vp in [("phone", {"width": 375, "height": 800}), ("desktop", {"width": 1280, "height": 900})]:
            ctx = browser.new_context(viewport=vp, accept_downloads=True)
            page = ctx.new_page()
            try:
                errs, xp, dp = play_all(page, tag, dl_dir)
                if errs:
                    failures.append(f"{tag}: console errors: {errs}")
                check_files(xp, dp)
                print(f"{tag}: full play-through OK, downloaded files verified")
            except Exception as e:
                failures.append(f"{tag}: {e!r}")
                page.screenshot(path=f"/tmp/fail-agent-{tag}.png", full_page=True)
            ctx.close()

        # class + teacher, double-click guard, phone overflow
        ctx = browser.new_context(viewport={"width": 375, "height": 800})
        page = ctx.new_page()
        errs = []
        page.on("pageerror", lambda e: errs.append(str(e)))
        page.goto(f"http://127.0.0.1:{PORT}/be-the-agent/index.html?mode=class&teacher=1")
        try:
            assert page.locator(".level-card[disabled]").count() == 0
            assert page.evaluate("getComputedStyle(document.documentElement).fontSize") == "18px"
            assert page.evaluate("document.documentElement.scrollWidth") <= 376
            page.locator('[data-level="L6"]').click()
            start(page)
            for i in range(8):
                page.get_by_role("button", name="✓ Approve").click()
                page.get_by_role("button", name="Next request" if i < 7 else "See results").click()
            first_opt(page)
            page.get_by_role("button", name="Finish level").dblclick()
            page.wait_for_timeout(300)
            n = page.get_by_text("Level 6 complete").count()
            assert n == 1, f"double-click gave {n} result blocks"
            print("class/teacher mode + double-click guard OK")
        except Exception as e:
            failures.append(f"class/teacher: {e!r}")
            page.screenshot(path="/tmp/fail-agent-class.png", full_page=True)
        if errs:
            failures.append(f"class/teacher errors: {errs}")
        ctx.close()

        # blocked storage
        ctx = browser.new_context()
        page = ctx.new_page()
        errs = []
        page.on("pageerror", lambda e: errs.append(str(e)))
        page.add_init_script("Object.defineProperty(window, 'localStorage', { get() { throw new Error('blocked'); } });")
        page.goto(f"http://127.0.0.1:{PORT}/be-the-agent/index.html")
        if page.locator(".level-card").count() != 6 or errs:
            failures.append(f"blocked storage: {errs}")
        else:
            print("blocked localStorage OK")
        ctx.close()
        browser.close()
    httpd.shutdown()
    if failures:
        print("FAILED:\n  " + "\n  ".join(failures))
        sys.exit(1)
    print("All play-through checks passed")

if __name__ == "__main__":
    main()
